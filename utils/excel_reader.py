from openpyxl import load_workbook


def get_excel_data(file_path):

    workbook = load_workbook(file_path)

    sheet = workbook["Sheet1"]

    data = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        data.append(row[0])

    return data